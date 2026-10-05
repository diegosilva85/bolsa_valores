"""Livro de operações em Excel e cálculo de posições, independente da interface."""
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path
import os
import re
import shutil
import tempfile
from uuid import uuid4
from .catalog import FX_SYMBOLS

CLASSES = ('Ações', 'FIIs', 'ETFs', 'BDRs', 'FIagro', 'FI-Infra', 'Outros', 'Ações EUA', 'ETFs EUA', 'Criptomoedas', 'Moedas')
LEGACY_HEADERS = ('ID', 'Data', 'Ticker', 'Nome', 'Classe', 'Operação', 'Quantidade', 'Preço unitário (R$)')
HEADERS = (*LEGACY_HEADERS[:7], 'Preço unitário', 'Moeda')


def currency_for(kind):
    return 'USD' if kind in ('Ações EUA', 'ETFs EUA', 'Criptomoedas') else 'BRL'


def symbol_for(ticker, kind):
    if kind == 'Moedas':
        return FX_SYMBOLS[ticker]
    return ticker if currency_for(kind) == 'USD' else ticker + '.SA'


def number(value):
    text = str(value).strip().replace('R$', '').replace(' ', '')
    if ',' in text:
        text = text.replace('.', '').replace(',', '.')
    try:
        result = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError('Informe um número válido, por exemplo 12,50.') from exc
    if not result.is_finite() or result <= 0:
        raise ValueError('Quantidade e preço devem ser positivos e finitos.')
    return result


def trade_date(value):
    if isinstance(value, datetime):
        value = value.date()
    if not isinstance(value, date):
        for fmt in ('%d/%m/%Y', '%Y-%m-%d'):
            try:
                value = datetime.strptime(str(value), fmt).date()
                break
            except ValueError:
                continue
        else:
            raise ValueError('Data inválida; use DD/MM/AAAA.')
    if value > date.today():
        raise ValueError('A operação não pode ter data futura.')
    return value


@dataclass(frozen=True)
class Trade:
    id: str
    day: date
    ticker: str
    name: str
    kind: str
    side: str
    quantity: Decimal
    price: Decimal

    @property
    def currency(self):
        return currency_for(self.kind)

    @classmethod
    def make(cls, day, ticker, name, kind, side, quantity, price, id=None):
        ticker = str(ticker).strip().upper().removesuffix('.SA')
        if kind == 'Moedas':
            ticker = ticker if ticker.endswith('/BRL') else ticker + '/BRL'
            valid = ticker in FX_SYMBOLS
        elif kind == 'Criptomoedas':
            ticker = ticker if ticker.endswith('-USD') else ticker + '-USD'
            valid = re.fullmatch(r'[A-Z0-9]{2,15}-USD', ticker)
        elif kind in ('Ações EUA', 'ETFs EUA'):
            ticker = ticker.replace('.', '-')
            valid = re.fullmatch(r'[A-Z]{1,8}(?:-[A-Z])?', ticker)
        else:
            valid = re.fullmatch(r'[A-Z]{4}[0-9]{1,2}', ticker)
        if not valid:
            raise ValueError('Ticker inválido para a classe selecionada.')
        if kind not in CLASSES or side not in ('Compra', 'Venda'):
            raise ValueError('Selecione uma classe e Compra ou Venda.')
        name = str(name or ticker).strip()
        if name.startswith(('=', '+', '-', '@')):
            raise ValueError('Nome de ativo inválido.')
        return cls(str(id or uuid4()), trade_date(day), ticker, name, kind, side, number(quantity), number(price))


@dataclass
class Position:
    ticker: str
    name: str
    kind: str
    quantity: Decimal = Decimal(0)
    cost: Decimal = Decimal(0)


def positions(trades):
    result, ids = {}, set()
    # Ordenação estável: operações no mesmo dia mantêm a ordem de lançamento.
    for t in sorted(trades, key=lambda t: t.day):
        if t.id in ids:
            raise ValueError('ID de operação repetido na planilha.')
        ids.add(t.id)
        p = result.setdefault(t.ticker, Position(t.ticker, t.name, t.kind))
        if p.kind != t.kind:
            raise ValueError(f'{t.ticker}: classe diferente entre operações.')
        if t.side == 'Compra':
            p.cost += t.quantity * t.price
            p.quantity += t.quantity
        else:
            if t.quantity > p.quantity:
                raise ValueError(f'{t.ticker}: venda em {t.day:%d/%m/%Y} excede o saldo de {p.quantity}.')
            p.cost -= p.cost / p.quantity * t.quantity
            p.quantity -= t.quantity
    return {key: p for key, p in result.items() if p.quantity > 0}


class ExcelPortfolio:
    def __init__(self, path):
        self.path = Path(path).resolve()
        self.trades = []
        self.fingerprint = None

    def digest(self):
        return sha256(self.path.read_bytes()).hexdigest() if self.path.exists() else None

    def load(self):
        from openpyxl import load_workbook
        before = self.digest()
        workbook = load_workbook(self.path, read_only=True, data_only=False)
        try:
            if 'Operações' not in workbook.sheetnames:
                raise ValueError('Arquivo inválido: falta a aba Operações.')
            rows = workbook['Operações'].iter_rows(values_only=True)
            header = tuple(next(rows, ()))
            if header not in (HEADERS, LEGACY_HEADERS):
                raise ValueError('Colunas incompatíveis. Abra um Excel criado pelo aplicativo.')
            trades = []
            for line, row in enumerate(rows, 2):
                if all(v is None for v in row):
                    continue
                try:
                    id, day, ticker, name, kind, side, quantity, price = row[:8]
                    currency = row[8] if header == HEADERS else 'BRL'
                    if currency != currency_for(kind):
                        raise ValueError('Moeda incompatível com a classe do ativo.')
                    if not id:
                        raise ValueError('ID ausente.')
                    trades.append(Trade.make(day, ticker, name, kind, side, quantity, price, id))
                except (ValueError, TypeError) as exc:
                    raise ValueError(f'Linha {line}: {exc}') from exc
            positions(trades)
        finally:
            workbook.close()
        if before != self.digest():
            raise ValueError('Arquivo alterado durante a leitura. Reabra a carteira.')
        self.trades, self.fingerprint = trades, before
        return self

    def write(self, trades):
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill
        positions(trades)
        if self.digest() != self.fingerprint:
            raise ValueError('O Excel foi alterado fora do app. Recarregue antes de salvar.')
        workbook = load_workbook(self.path) if self.path.exists() else Workbook()
        temporary = None
        try:
            if 'Operações' in workbook.sheetnames:
                sheet = workbook['Operações']
                sheet.delete_rows(1, sheet.max_row)
            else:
                sheet = workbook.active
                sheet.title = 'Operações'
            sheet.append(HEADERS)
            for t in trades:
                sheet.append([t.id, t.day, t.ticker, t.name, t.kind, t.side, float(t.quantity), float(t.price), t.currency])
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = f'A1:I{max(1, sheet.max_row)}'
            for cell in sheet[1]:
                cell.font = Font(color='FFFFFF', bold=True)
                cell.fill = PatternFill('solid', fgColor='15324F')
            for column, width in zip('ABCDEFGHI', (39, 15, 14, 40, 16, 14, 19, 23, 12)):
                sheet.column_dimensions[column].width = width
            for row in sheet.iter_rows(min_row=2):
                row[1].number_format = 'dd/mm/yyyy'
                row[6].number_format = '#,##0.########'
                row[7].number_format = ('"US$" ' if row[8].value == 'USD' else '"R$" ') + '#,##0.00######'
                for cell in row:
                    cell.font = Font(color='0000FF')
            self.path.parent.mkdir(parents=True, exist_ok=True)
            descriptor, temporary = tempfile.mkstemp(suffix='.xlsx', dir=self.path.parent)
            os.close(descriptor)
            workbook.save(temporary)
            if self.digest() != self.fingerprint:
                raise ValueError('Arquivo alterado durante a gravação. Recarregue a carteira.')
            if self.path.exists():
                shutil.copy2(self.path, self.path.with_suffix('.backup.xlsx'))
            os.replace(temporary, self.path)
            self.trades, self.fingerprint = list(trades), self.digest()
        finally:
            workbook.close()
            if temporary and os.path.exists(temporary):
                os.unlink(temporary)

    def add(self, trade):
        self.write([*self.trades, trade])
