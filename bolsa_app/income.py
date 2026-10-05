"""Proventos por data de pagamento, com posição histórica e cache auditável."""
from dataclasses import asdict, dataclass, replace
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from hashlib import sha256
import json
import os
import re
import unicodedata
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

TYPES = ('Dividendos', 'JCP', 'Rendimentos FII', 'Amortizações', 'Outros rendimentos')
INCOME_HEADERS = ('ID', 'Ticker', 'Classe', 'Tipo', 'Data com', 'Pagamento', 'Valor unitário',
                  'Quantidade com direito', 'Valor bruto', 'Moeda', 'Fonte', 'Observação', 'Revisar')
PERIOD_HEADERS = ('Ano', 'Gerado em', 'Histórico da carteira', 'Situação', 'Avisos')


def history_key(trades, events):
    content = json.dumps([[asdict(t) for t in trades], [asdict(e) for e in events]],
                         default=str, sort_keys=True, ensure_ascii=False)
    return sha256(content.encode()).hexdigest()


def parse_day(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def decimal_value(value, positive=False):
    try:
        result = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError('Valor numérico inválido no provento.') from exc
    if not result.is_finite() or result < 0 or (positive and result == 0):
        raise ValueError('Valor do provento deve ser finito e não negativo.')
    return result


@dataclass(frozen=True)
class Income:
    id: str
    ticker: str
    kind: str
    type: str
    record_day: date
    pay_day: date
    unit: Decimal
    quantity: Decimal
    gross: Decimal
    currency: str
    source: str
    note: str
    review: bool

    @classmethod
    def from_row(cls, row):
        if len(row) != len(INCOME_HEADERS):
            raise ValueError('Número de colunas de Proventos inválido.')
        id, ticker, kind, type, record, pay, unit, qty, gross, currency, source, note, review = row
        from .portfolio import CLASSES
        if not id or not re.fullmatch(r'[A-Z0-9.^/-]{1,25}', str(ticker)) or kind not in CLASSES:
            raise ValueError('Identificação inválida em Proventos.')
        if type not in TYPES or currency not in ('BRL', 'USD') or not isinstance(review, bool):
            raise ValueError('Tipo, moeda ou revisão inválida em Proventos.')
        unit, qty, gross = decimal_value(unit), decimal_value(qty, True), decimal_value(gross)
        if gross != (unit * qty).quantize(Decimal('.01'), rounding=ROUND_HALF_UP):
            raise ValueError('Total de Proventos diferente de quantidade × valor unitário.')
        record, pay = parse_day(record), parse_day(pay)
        if record > pay:
            raise ValueError('Data com posterior ao pagamento.')
        return cls(str(id), ticker, kind, type, record, pay, unit, qty, gross, currency,
                   str(source or ''), str(note or ''), review)


class BrapiIncomeProvider:
    def __init__(self, token=None):
        self.token = os.getenv('BRAPI_TOKEN', '') if token is None else token

    def fetch(self, ticker, kind, year):
        if kind in ('FIagro', 'FI-Infra'):
            raise ValueError('Esta integração cobre ações/BDRs/ETFs e FIIs; este tipo de fundo requer outra rota da fonte.')
        if not re.fullmatch(r'[A-Z]{4}\d{1,2}', ticker):
            raise ValueError('Consulta automática disponível somente para ativos B3.')
        fund = kind == 'FIIs'
        route = 'fii' if fund else 'stocks'
        params = dict(symbols=ticker, startDate=f'{year}-01-01', endDate=f'{year}-12-31', sortOrder='asc')
        if not fund:
            # rate pode estar ajustado a splits posteriores. Não multiplicar pela
            # posição histórica sem o valor original (rawRate, plano Pro).
            params['includeRaw'] = 'true'
            params['sortBy'] = 'paymentDate'
        headers = {'User-Agent': 'BolsaBrasil/1.0'}
        if self.token:
            headers['Authorization'] = 'Bearer ' + self.token
        request = Request(f'https://brapi.dev/api/v2/{route}/dividends?' + urlencode(params), headers=headers)
        try:
            with urlopen(request, timeout=20) as response:
                payload = json.load(response)
        except HTTPError as exc:
            if exc.code in (401, 403):
                raise ValueError('Acesso negado: configure token/plano brapi (ações exigem rawRate/includeRaw).') from None
            raise ValueError(f'brapi HTTP {exc.code}; tente gerar novamente mais tarde.') from None
        except (URLError, TimeoutError, OSError, ValueError):
            raise ValueError('Não foi possível consultar a brapi; verifique conexão e tente novamente.') from None
        if fund:
            rows = payload.get('dividends')
            if not isinstance(rows, list):
                raise ValueError('Resposta de rendimentos inválida.')
            if any(r.get('symbol') != ticker for r in rows):
                raise ValueError('A fonte retornou outro ticker; confira a mudança de ativo.')
        else:
            results = payload.get('results')
            if not isinstance(results, list) or len(results) != 1:
                raise ValueError('Resposta de dividendos incompleta.')
            item = results[0]
            if item.get('symbol') != ticker or item.get('requestedSymbol') != ticker or item.get('changed'):
                raise ValueError('A fonte redirecionou o ticker; histórico requer conferência manual.')
            rows = item.get('data', {}).get('cashDividends')
            if not isinstance(rows, list):
                raise ValueError('Histórico de proventos indisponível.')
        return rows


def income_type(label, kind):
    label = ''.join(c for c in unicodedata.normalize('NFD', str(label).upper()) if not unicodedata.combining(c))
    if 'AMORTIZA' in label:
        return 'Amortizações'
    if label == 'JCP' or 'JUROS SOBRE' in label:
        return 'JCP'
    if 'DIVIDEND' in label:
        return 'Dividendos'
    if 'RENDIMENTO' in label:
        return 'Rendimentos FII' if kind in ('FIIs', 'FIagro', 'FI-Infra') else 'Outros rendimentos'
    raise ValueError(f'Tipo de provento não suportado: {label[:60]}')


def generate_income(trades, events, year, provider, today=None, cancel=None):
    from .portfolio import positions
    today = today or date.today()
    if not 1900 <= year <= today.year:
        raise ValueError('Selecione um ano entre 1900 e o ano atual.')
    positions(trades, events)  # não gerar a partir de histórico inconsistente
    end = date(year, 12, 31)
    assets = {t.ticker: t.kind for t in trades if t.day <= end}
    for e in events:
        if e.day <= end:
            assets.setdefault(e.source, e.kind)
            assets.setdefault(e.target, e.target_kind)
    rows, warnings, successful = [], [], []
    snapshots = {}
    seen = set()
    for ticker, kind in sorted(assets.items()):
        if cancel is not None and cancel.is_set():
            raise ValueError('Consulta cancelada.')
        if kind in ('Criptomoedas', 'Moedas'):
            continue
        try:
            fetched = provider.fetch(ticker, kind, year)
        except Exception as exc:
            warnings.append(f'{ticker}: {exc}')
            continue
        successful.append(ticker)
        for event in fetched:
            try:
                pay = parse_day(event.get('paymentDate'))
                if pay.year != year:
                    continue
                record = parse_day(event.get('lastDatePrior'))
                if record > pay:
                    raise ValueError('Data com posterior ao pagamento.')
                if record > today:
                    raise ValueError('Data de direito futura: posição ainda não conhecida.')
                if record not in snapshots:
                    snapshots[record] = positions([t for t in trades if t.day <= record],
                                                 [e for e in events if e.day <= record])
                position = snapshots[record].get(ticker)
                if position is None:
                    continue
                category = income_type(event.get('label'), kind)
                raw = event.get('rate') if kind in ('FIIs', 'FIagro', 'FI-Infra') else event.get('rawRate')
                if raw is None:
                    raise ValueError('Valor unitário original ausente (rawRate); não usar valor ajustado.')
                unit = decimal_value(raw, True)
                # Identidade sem o valor: duplicatas idênticas são ignoradas;
                # valores conflitantes não são somados como pagamentos distintos.
                identity = json.dumps([ticker, category, str(record), str(pay), event.get('relatedTo') or '',
                                       event.get('isinCode') or ''], ensure_ascii=False)
                id = sha256(identity.encode()).hexdigest()
                if id in seen:
                    existing = next(r for r in rows if r.id == id)
                    if existing.unit != unit:
                        rows[rows.index(existing)] = replace(existing, review=True, note='Valores conflitantes na fonte; conferir retificação.')
                        raise ValueError('Versões conflitantes do mesmo provento; conferir retificação.')
                    continue
                seen.add(id)
                review = event.get('verified') is not True or bool(event.get('paymentDateIsDeadline'))
                note = 'Conferir com extrato; líquido e retenções não calculados.'
                if review:
                    note += ' Fonte não verificada ou informa somente prazo de pagamento.'
                rows.append(Income(id, ticker, position.kind, category, record, pay, unit, position.quantity,
                    (unit * position.quantity).quantize(Decimal('.01'), rounding=ROUND_HALF_UP), 'BRL',
                    'brapi', note, review))
            except (ValueError, TypeError, KeyError) as exc:
                warnings.append(f'{ticker}: evento ignorado/requer revisão: {exc}')
    # Amortizações do livro corporativo não são somadas automaticamente:
    # poderiam duplicar a fonte e a data efetiva não é necessariamente pagamento.
    return rows, {'year': year, 'generated': datetime.now().isoformat(timespec='seconds'),
                  'history': history_key(trades, events), 'status': 'Parcial' if warnings else 'Consultado',
                  'warnings': '\n'.join(warnings)}, successful


def monthly_totals(rows, year, today=None):
    today = today or date.today()
    result = {m: {kind: Decimal(0) for kind in TYPES} for m in range(1, 13)}
    for row in rows:
        if row.pay_day.year == year and row.pay_day <= today and not row.review and row.currency == 'BRL':
            result[row.pay_day.month][row.type] += row.gross
    return result


def read_income(workbook):
    rows, periods = [], {}
    if 'Proventos' in workbook.sheetnames:
        iterator = workbook['Proventos'].iter_rows(values_only=True)
        if tuple(next(iterator, ())) != INCOME_HEADERS:
            raise ValueError('Colunas da aba Proventos incompatíveis.')
        for line, values in enumerate(iterator, 2):
            if all(v is None for v in values):
                continue
            try:
                rows.append(Income.from_row(values))
            except (ValueError, TypeError) as exc:
                raise ValueError(f'Proventos, linha {line}: {exc}') from exc
        if len({r.id for r in rows}) != len(rows):
            raise ValueError('IDs duplicados na aba Proventos.')
    if 'PeriodosProventos' in workbook.sheetnames:
        iterator = workbook['PeriodosProventos'].iter_rows(values_only=True)
        if tuple(next(iterator, ())) != PERIOD_HEADERS:
            raise ValueError('Colunas dos períodos de proventos incompatíveis.')
        for values in iterator:
            if all(v is None for v in values):
                continue
            year, generated, history, status, warnings = values
            year = int(year)
            if not 1900 <= year <= date.today().year or year in periods or status not in ('Consultado', 'Parcial'):
                raise ValueError('Período de proventos inválido ou repetido.')
            datetime.fromisoformat(str(generated))
            periods[year] = dict(year=year, generated=str(generated), history=str(history), status=status, warnings=str(warnings or ''))
    return rows, periods


def write_income(workbook, rows, periods):
    from openpyxl.styles import Font, PatternFill
    for name, headers, values in [
        ('Proventos', INCOME_HEADERS, [[str(v) if isinstance(v, Decimal) else v for v in asdict(r).values()] for r in rows]),
        ('PeriodosProventos', PERIOD_HEADERS, [[p[k] for k in ('year', 'generated', 'history', 'status', 'warnings')] for _, p in sorted(periods.items())])]:
        sheet = workbook[name] if name in workbook.sheetnames else workbook.create_sheet(name)
        sheet.delete_rows(1, sheet.max_row)
        sheet.append(headers)
        for row in values:
            sheet.append(row)
        # Strings externas são texto literal, nunca fórmulas do Excel.
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                elif isinstance(cell.value, date):
                    cell.number_format = 'dd/mm/yyyy'
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='15324F')
            sheet.column_dimensions[cell.column_letter].width = 24
