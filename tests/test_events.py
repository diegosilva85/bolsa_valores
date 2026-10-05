import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from bolsa_app.portfolio import CorporateEvent, ExcelPortfolio, Trade, positions


class EventsTests(unittest.TestCase):
    def test_liquidation_calculator_uses_historical_adjusted_position(self):
        from bolsa_app.portfolio import calculate_liquidation
        split = self.event('Desdobramento', factor=8)
        future_sale = Trade.make('05/01/2020', 'ABCD3', '', 'Ações', 'Venda', 80, 10)
        result = calculate_liquidation([self.buy(), future_sale], [split], '03/01/2020',
            'Antes', 'ABCD3', '0,8', '8,20', '0,56', '2,00', '0')
        self.assertEqual(result['source_quantity'], 80)
        self.assertEqual(result['received_quantity'], 64)
        self.assertEqual(result['received_price'], Decimal('10.25'))
        self.assertEqual(result['cash'], Decimal('42.80'))

    def test_calculator_fraction_timing_and_invalid_inputs(self):
        from bolsa_app.portfolio import calculate_liquidation
        def calc(**kwargs):
            args = dict(day='01/01/2020', timing='Depois', source='ABCD3', factor='.83',
                        asset_value_per_old='8.3', cash_per_old='.56', deductions='0', fraction_cash='1.23')
            args.update(kwargs)
            return calculate_liquidation([self.buy()], [], **args)
        result = calc()
        self.assertEqual(result['fraction'], Decimal('.3'))
        self.assertEqual(result['cash'], Decimal('6.83'))
        for kwargs in (dict(timing='Antes'), dict(deductions=100), dict(cash_per_old=''),
                       dict(cash_per_old='NaN'), dict(fraction_cash=-1), dict(factor=0)):
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                calc(**kwargs)

    def test_excel_class_correction_reconciles_events(self):
        from openpyxl import load_workbook
        for ticker in ('BCFF11', 'BBPO11'):
            with self.subTest(ticker=ticker), tempfile.TemporaryDirectory() as folder:
                trade = Trade.make('01/01/2018', ticker, '', 'Ações', 'Compra', 10, 80)
                event = CorporateEvent.make('29/11/2023', 'Desdobramento', ticker, 'Ações', factor=8)
                store = ExcelPortfolio(Path(folder) / 'test.xlsx')
                store.write([trade], [event])
                workbook = load_workbook(store.path)
                workbook['Operações']['E2'] = 'FIIs'
                workbook.save(store.path)
                workbook.close()
                before = store.path.read_bytes()
                loaded = ExcelPortfolio(store.path).load()
                self.assertEqual(store.path.read_bytes(), before)
                self.assertEqual((loaded.events[0].kind, loaded.events[0].target_kind), ('FIIs', 'FIIs'))
                p = positions(loaded.trades, loaded.events)[ticker]
                self.assertEqual((p.kind, p.quantity, p.cost), ('FIIs', 80, 800))
                loaded.write(loaded.trades)
                self.assertEqual(ExcelPortfolio(store.path).load().events[0].kind, 'FIIs')

    def test_class_correction_does_not_hide_missing_position(self):
        trade = Trade.make('03/01/2020', 'ABCD3', '', 'FIIs', 'Compra', 10, 20)
        with self.assertRaisesRegex(ValueError, 'sem saldo na data'):
            positions([trade], [self.event('Desdobramento', factor=8)])

    def test_liquidation_distinct_cost_cash_fraction_and_sale(self):
        event = self.event('Liquidação com entrega de cotas', target='EFGH3', factor='.83',
                           received_quantity=8, received_price=12, cash='5.61')
        ledger = []
        existing = Trade.make('01/01/2020', 'EFGH3', '', 'Ações', 'Compra', 2, 10)
        sell = Trade.make('03/01/2020', 'EFGH3', '', 'Ações', 'Venda', 5, 15)
        result = positions([self.buy(), existing, sell], [event], ledger)
        self.assertNotIn('ABCD3', result)
        self.assertEqual((result['EFGH3'].quantity, result['EFGH3'].cost), (5, 58))
        self.assertEqual(ledger[0]['fraction'], Decimal('.3'))
        self.assertEqual(ledger[0]['cash'], Decimal('5.61'))
        self.assertEqual(ledger[0]['old_cost'], 200)
        with tempfile.TemporaryDirectory() as folder:
            store = ExcelPortfolio(Path(folder) / 'liquidation.xlsx')
            store.write([self.buy()], [event])
            self.assertEqual(ExcelPortfolio(store.path).load().events, [event])

    def test_liquidation_validation(self):
        with self.assertRaises(ValueError):
            self.event('Liquidação com entrega de cotas', target='EFGH3')
        for qty in (7, 9):
            event = self.event('Liquidação com entrega de cotas', target='EFGH3', factor='.83', received_quantity=qty, received_price=12)
            with self.assertRaises(ValueError):
                positions([self.buy()], [event])

    def test_legacy_event_sheet(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            store = ExcelPortfolio(Path(folder) / 'legacy.xlsx')
            event = self.event('Desdobramento', factor=8)
            store.write([self.buy()], [event])
            workbook = load_workbook(store.path)
            workbook['Eventos'].delete_cols(13, 3)
            workbook.save(store.path)
            workbook.close()
            loaded = ExcelPortfolio(store.path).load()
            self.assertEqual(loaded.events, [event])
            loaded.write(loaded.trades)
            self.assertEqual(ExcelPortfolio(store.path).load().events, [event])

    def buy(self):
        return Trade.make('01/01/2020', 'ABCD3', 'Original', 'Ações', 'Compra', 10, 20)

    def event(self, type, **kwargs):
        return CorporateEvent.make('02/01/2020', type, 'ABCD3', 'Ações', **kwargs)

    def test_split_and_sale(self):
        event = self.event('Desdobramento', factor=2)
        sell = Trade.make('03/01/2020', 'ABCD3', '', 'Ações', 'Venda', 15, 12)
        p = positions([self.buy(), sell], [event])['ABCD3']
        self.assertEqual((p.quantity, p.cost), (5, 50))
        with self.assertRaises(ValueError):
            positions([self.buy(), sell])

    def test_reverse_fraction(self):
        p = positions([self.buy()], [self.event('Grupamento', factor='0.03')])['ABCD3']
        self.assertEqual((p.quantity, p.cost), (Decimal('.30'), 200))

    def test_merger_and_spinoff(self):
        for type, expected in [('Incorporação', 200), ('Cisão', 50)]:
            result = positions([self.buy()], [self.event(type, target='EFGH3', factor=2, cost_percent=25)])
            self.assertEqual(result['EFGH3'].quantity, 20)
            self.assertEqual(result['EFGH3'].cost, expected)
            self.assertEqual(sum(p.cost for p in result.values()), 200)

    def test_rename_bonus_amortization(self):
        result = positions([self.buy()], [self.event('Troca de ticker/nome', target='EFGH3', name='Novo')])
        self.assertNotIn('ABCD3', result)
        self.assertEqual(result['EFGH3'].name, 'Novo')
        p = positions([self.buy()], [self.event('Bonificação', factor='.1', amount=5)])['ABCD3']
        self.assertEqual((p.quantity, p.cost), (11, 205))
        p = positions([self.buy()], [self.event('Amortização', amount=2)])['ABCD3']
        self.assertEqual(p.cost, 180)

    def test_validation_and_same_day(self):
        event = self.event('Desdobramento', factor=2)
        with self.assertRaises(ValueError):
            positions([self.buy()], [event, self.event('Desdobramento', factor=2)])
        with self.assertRaises(ValueError):
            positions([self.buy()], [self.event('Amortização', amount=21)])
        with self.assertRaises(ValueError):
            positions([], [event])
        event = CorporateEvent.make('01/01/2020', 'Desdobramento', 'ABCD3', 'Ações', factor=2, timing='Depois')
        self.assertEqual(positions([self.buy()], [event])['ABCD3'].quantity, 20)

    def test_excel_atomic_roundtrip(self):
        with tempfile.TemporaryDirectory() as folder:
            store = ExcelPortfolio(Path(folder) / 'portfolio.xlsx')
            event = self.event('Troca de ticker/nome', target='EFGH3', name='Novo')
            store.write([self.buy()], [event])
            loaded = ExcelPortfolio(store.path).load()
            self.assertEqual(loaded.events, [event])
            self.assertEqual(loaded.trades, store.trades)
            before = store.path.read_bytes()
            with self.assertRaises(ValueError):
                store.write([], [event])
            self.assertEqual(store.path.read_bytes(), before)
            store.write(store.trades)
            self.assertEqual(ExcelPortfolio(store.path).load().events, [event])
